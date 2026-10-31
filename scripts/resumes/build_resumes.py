"""Build the two résumé DOCX files from the same content used by /resume.
Run with the Python runtime returned by the Codex workspace dependency loader.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
DATA = json.loads((ROOT / 'lib/resumes.json').read_text())
OUT = ROOT / 'private/resumes'
OUT.mkdir(parents=True, exist_ok=True)


def link(p, label, url, size=9.5):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); pr.append(color)
    font = OxmlElement('w:rFonts'); font.set(qn('w:ascii'), 'Arial'); font.set(qn('w:hAnsi'), 'Arial'); pr.append(font)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); pr.append(sz)
    r.append(pr); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)


def heading(doc, text):
    p = doc.add_paragraph(text, 'Heading 1')
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text):
    p = doc.add_paragraph(text, 'List Bullet')
    p.paragraph_format.widow_control = True
    return p


for version in DATA['versions']:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.2677); section.page_height = Inches(11.6929)
    section.top_margin = Inches(.52); section.bottom_margin = Inches(.52)
    section.left_margin = Inches(.62); section.right_margin = Inches(.62)
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.04
    for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet']:
        style = doc.styles[name]; style.font.name = 'Arial'; style.font.color.rgb = RGBColor(0, 0, 0)
    title = doc.styles['Title']
    for border in list(title.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    title.font.size = Pt(25); title.font.bold = True
    title.paragraph_format.space_after = Pt(2)
    h = doc.styles['Heading 1']; h.font.size = Pt(11); h.font.bold = True
    h.paragraph_format.space_before = Pt(9); h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    lb = doc.styles['List Bullet']; lb.font.size = Pt(10.5)
    lb.paragraph_format.space_after = Pt(3); lb.paragraph_format.line_spacing = 1.04
    lb.paragraph_format.left_indent = Inches(.14); lb.paragraph_format.first_line_indent = Inches(-.14)
    doc.core_properties.author = DATA['name']; doc.core_properties.title = f"{DATA['name']} {version['title']} Resume"
    doc.core_properties.subject = version['audience']; doc.core_properties.keywords = ''
    doc.add_paragraph(DATA['name'], 'Title')
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    run = p.add_run(version['title']); run.bold = True; run.font.size = Pt(12)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(DATA['location'] + '  |  '); r.font.size = Pt(9.5)
    link(p, DATA['email'], 'mailto:' + DATA['email']); p.add_run('  |  ')
    link(p, 'kimtsok.vercel.app', DATA['website'])
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    link(p, 'github.com/Kim-Tsok', DATA['github']); p.add_run('  |  ')
    link(p, 'linkedin.com/in/kim-tsok-5414a8418', DATA['linkedin'])
    doc.add_paragraph(version['summary'])
    heading(doc, 'Skills')
    for text in version['skills']:
        p = doc.add_paragraph(text); p.paragraph_format.space_after = Pt(2)
    heading(doc, 'Experience')
    for job in version['experience']:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1); p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.02), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(job['role']).bold = True
        date = p.add_run('\t' + job['dates']); date.font.size = Pt(9)
        p = doc.add_paragraph(job['organization']); p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        for r in p.runs: r.font.size = Pt(9.5)
        for text in job['bullets']: bullet(doc, text)
    heading(doc, 'Selected projects')
    for project in version['projects']:
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
        p.add_run(project['name']).bold = True; p.add_run('  |  ')
        link(p, project['url'].removeprefix('https://').rstrip('/'), project['url'])
        p = doc.add_paragraph(project['stack']); p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        for r in p.runs: r.font.size = Pt(9.5)
        for text in project['bullets']: bullet(doc, text)
    heading(doc, 'Education')
    p = doc.add_paragraph(DATA['education']); p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(1)
    doc.add_paragraph(DATA['educationDates'])
    heading(doc, 'Certification and award')
    for text in DATA['credentials']:
        p = doc.add_paragraph(text); p.paragraph_format.space_after = Pt(1)
    path = OUT / (version['filename'] + '.docx')
    doc.save(path)
    print(path)
